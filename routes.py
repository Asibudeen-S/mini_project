import csv
import io
import os
import uuid
from datetime import date, datetime, timezone
from functools import wraps
from urllib.parse import urljoin, urlparse

from flask import (Blueprint, Response, abort, current_app, flash, make_response,
                   redirect, render_template, request, send_from_directory,
                   url_for)
from flask_login import current_user, login_user, logout_user, login_required
from sqlalchemy import func, or_
from werkzeug.utils import secure_filename

from extensions import db
from forms import (AttendanceSelectorForm, DepartmentForm, LeaveRequestForm,
                   AdminCreateUserForm, LoginForm, ProfileForm, ReviewLeaveForm,
                   SettingsForm, SubjectForm, UserAdminForm)
from models import (AttendanceItem, AttendanceRecord, AuditLog, Department,
                    Enrollment, LeaveRequest, Subject, SystemSetting, User, utcnow)


bp = Blueprint("main", __name__)
ALLOWED_UPLOAD_EXTENSIONS = {"pdf", "png", "jpg", "jpeg", "doc", "docx"}
ATTENDANCE_STATUSES = ("present", "absent", "late", "excused")
MAX_STUDENTS_PER_DEPARTMENT = 100


def role_required(*roles):
    def decorator(view):
        @wraps(view)
        @login_required
        def wrapped(*args, **kwargs):
            if current_user.role not in roles:
                abort(403)
            return view(*args, **kwargs)
        return wrapped
    return decorator


def audit(action, entity_type, entity_id=None, details=""):
    db.session.add(AuditLog(actor_id=current_user.id if current_user.is_authenticated else None,
                            action=action, entity_type=entity_type,
                            entity_id=str(entity_id) if entity_id is not None else None,
                            details=details[:4000]))


def setting_value(key, default):
    setting = db.session.get(SystemSetting, key)
    return setting.value if setting else default


def safe_spreadsheet_cell(value):
    cell = "" if value is None else str(value)
    return "'" + cell if cell.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) else cell


def attendance_percentage(items):
    countable = [item for item in items if item.status != "excused"]
    if not countable:
        return None
    attended = sum(item.status in ("present", "late") for item in countable)
    return round(attended * 100 / len(countable), 1)


def user_subjects(user):
    if user.role == "admin":
        return Subject.query.order_by(Subject.name).all()
    if user.role == "teacher":
        if user.department_id:
            return Subject.query.filter_by(department_id=user.department_id).order_by(Subject.name).all()
        return Subject.query.filter_by(teacher_id=user.id).order_by(Subject.name).all()
    return (Subject.query.join(Enrollment).filter(
        Enrollment.student_id == user.id, Subject.study_year == user.study_year)
            .order_by(Subject.name).all())


def teacher_can_review(user, student_id):
    if user.department_id and User.query.filter_by(id=student_id, department_id=user.department_id, role="student").first():
        return True
    return (Enrollment.query.join(Subject, Enrollment.subject_id == Subject.id)
            .filter(Enrollment.student_id == student_id, Subject.teacher_id == user.id).first() is not None)


def valid_batch_student_ids(student_ids, department_id, study_year):
    student_ids = set(student_ids or [])
    if not student_ids:
        return set()
    return {row[0] for row in db.session.query(User.id).filter(
        User.id.in_(student_ids), User.role == "student", User.status == "active",
        User.department_id == department_id, User.study_year == study_year).all()}


def get_threshold():
    try:
        return float(setting_value("attendance_threshold", "75"))
    except (TypeError, ValueError):
        return 75.0


def report_query():
    query = AttendanceItem.query.join(AttendanceRecord).join(Subject).join(User, AttendanceItem.student_id == User.id)
    if current_user.role == "teacher":
        query = (query.filter(Subject.department_id == current_user.department_id)
                 if current_user.department_id else query.filter(Subject.teacher_id == current_user.id))
    from_date = parse_date(request.args.get("from"))
    to_date = parse_date(request.args.get("to"))
    department_id = parse_int(request.args.get("department"))
    subject_id = parse_int(request.args.get("subject"))
    user_id = parse_int(request.args.get("user"))
    if from_date:
        query = query.filter(AttendanceRecord.date >= from_date)
    if to_date:
        query = query.filter(AttendanceRecord.date <= to_date)
    if department_id:
        query = query.filter(Subject.department_id == department_id)
    if subject_id:
        query = query.filter(AttendanceRecord.subject_id == subject_id)
    if user_id:
        query = query.filter(AttendanceItem.student_id == user_id)
    return query.order_by(AttendanceRecord.date.desc(), Subject.name, User.name), {
        "from": request.args.get("from", ""), "to": request.args.get("to", ""),
        "department": department_id, "subject": subject_id, "user": user_id,
    }


def parse_date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date() if value else None
    except ValueError:
        return None


def parse_int(value):
    try:
        return int(value) if value else None
    except (TypeError, ValueError):
        return None


def safe_next_url(target):
    if not target:
        return None
    ref = urlparse(request.host_url)
    test = urlparse(urljoin(request.host_url, target))
    return target if test.scheme in {"http", "https"} and ref.netloc == test.netloc else None


@bp.route("/")
def index():
    if current_user.is_authenticated:
        return redirect(url_for("main.dashboard"))
    return redirect(url_for("main.login"))


@bp.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("main.dashboard"))
    form = LoginForm()
    if form.validate_on_submit():
        user = User.query.filter(func.lower(User.email) == form.email.data.strip().lower()).first()
        if user and user.check_password(form.password.data):
            if user.status != "active":
                flash("Your account is awaiting administrator approval or has been deactivated.", "warning")
            else:
                login_user(user, remember=False, fresh=True)
                flash(f"Welcome back, {user.name.split()[0]}.", "success")
                return redirect(safe_next_url(request.args.get("next")) or url_for("main.dashboard"))
        else:
            flash("Email or password was not recognized.", "danger")
    return render_template("login.html", form=form)


@bp.route("/register", methods=["GET", "POST"])
def register():
    flash("Student accounts are created by the administrator. Contact the administrator for an account.", "info")
    return redirect(url_for("main.login"))


@bp.route("/logout", methods=["POST"])
@login_required
def logout():
    logout_user()
    flash("You have been signed out.", "info")
    return redirect(url_for("main.login"))


@bp.route("/dashboard")
@login_required
def dashboard():
    threshold = get_threshold()
    if current_user.role == "admin":
        stats = {
            "users": User.query.filter_by(status="active").count(),
            "pending": User.query.filter_by(status="pending").count(),
            "classes": Subject.query.count(),
            "sessions": AttendanceRecord.query.count(),
        }
        recent_audits = AuditLog.query.order_by(AuditLog.created_at.desc()).limit(8).all()
        return render_template("dashboard.html", role="admin", stats=stats,
                               recent_audits=recent_audits, threshold=threshold)
    if current_user.role == "teacher":
        subjects = user_subjects(current_user)
        assigned_ids = [subject.id for subject in subjects]
        student_count = (db.session.query(func.count(func.distinct(Enrollment.student_id)))
                         .filter(Enrollment.subject_id.in_(assigned_ids)).scalar() if assigned_ids else 0)
        recent_records = (AttendanceRecord.query.filter(AttendanceRecord.subject_id.in_(assigned_ids))
                          .order_by(AttendanceRecord.date.desc()).limit(8).all() if assigned_ids else [])
        low_attendance = []
        for subject in subjects:
            for enrollment in subject.enrollments:
                items = (AttendanceItem.query.join(AttendanceRecord)
                         .filter(AttendanceRecord.subject_id == subject.id,
                                 AttendanceItem.student_id == enrollment.student_id).all())
                percentage = attendance_percentage(items)
                if percentage is not None and percentage < threshold:
                    low_attendance.append({"student": enrollment.student, "subject": subject, "percentage": percentage})
        low_attendance.sort(key=lambda row: row["percentage"])
        return render_template("dashboard.html", role="teacher", subjects=subjects,
                               student_count=student_count, recent_records=recent_records,
                               low_attendance=low_attendance[:10], threshold=threshold)

    items = (AttendanceItem.query.join(AttendanceRecord).filter(AttendanceItem.student_id == current_user.id)
             .order_by(AttendanceRecord.date.desc()).all())
    countable = [item for item in items if item.status != "excused"]
    attended = sum(item.status in ("present", "late") for item in countable)
    percentage = round(attended * 100 / len(countable), 1) if countable else None
    breakdown = {status: sum(item.status == status for item in items) for status in ATTENDANCE_STATUSES}
    subjects = user_subjects(current_user)
    subject_stats = []
    for subject in subjects:
        subject_items = [item for item in items if item.record.subject_id == subject.id]
        subject_stats.append({"subject": subject, "percentage": attendance_percentage(subject_items), "sessions": len(subject_items)})
    requests = LeaveRequest.query.filter_by(student_id=current_user.id).order_by(LeaveRequest.created_at.desc()).limit(5).all()
    return render_template("dashboard.html", role="student", percentage=percentage,
                           attended=attended, total=len(countable), breakdown=breakdown,
                           subject_stats=subject_stats, recent_items=items[:8],
                           leave_requests=requests, threshold=threshold,
                           schedule=sorted(subjects, key=lambda subject: (
                               ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday").index(subject.weekday)
                               if subject.weekday in ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday") else 7,
                               subject.start_time or "")))


@bp.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    form = ProfileForm(obj=current_user)
    if form.validate_on_submit():
        email = form.email.data.strip().lower()
        other = User.query.filter(func.lower(User.email) == email, User.id != current_user.id).first()
        if other:
            flash("That email is already in use.", "danger")
        elif form.new_password.data and not current_user.check_password(form.current_password.data or ""):
            flash("Enter your current password to set a new password.", "danger")
        else:
            current_user.name = form.name.data.strip()
            current_user.email = email
            if form.new_password.data:
                current_user.set_password(form.new_password.data)
            audit("profile_updated", "user", current_user.id)
            db.session.commit()
            flash("Profile updated.", "success")
            return redirect(url_for("main.profile"))
    return render_template("profile.html", form=form)


@bp.route("/admin/departments", methods=["GET", "POST"])
@role_required("admin")
def departments():
    form = DepartmentForm()
    if form.validate_on_submit():
        name, code = form.name.data.strip(), form.code.data.strip().upper()
        if Department.query.filter(or_(func.lower(Department.name) == name.lower(), func.lower(Department.code) == code.lower())).first():
            flash("A department with that name or code already exists.", "danger")
        else:
            department = Department(name=name, code=code)
            db.session.add(department)
            db.session.flush()
            audit("department_created", "department", department.id, code)
            db.session.commit()
            flash("Department created.", "success")
            return redirect(url_for("main.departments"))
    rows = Department.query.order_by(Department.name).all()
    student_counts = {
        department.id: User.query.filter_by(role="student", department_id=department.id)
        .filter(User.status.in_(["active", "pending"])).count()
        for department in rows
    }
    return render_template("departments.html", form=form, departments=rows,
                           student_counts=student_counts)


@bp.route("/admin/departments/<int:department_id>/edit", methods=["GET", "POST"])
@role_required("admin")
def edit_department(department_id):
    department = db.get_or_404(Department, department_id)
    form = DepartmentForm(obj=department)
    if form.validate_on_submit():
        name, code = form.name.data.strip(), form.code.data.strip().upper()
        duplicate = Department.query.filter(
            or_(func.lower(Department.name) == name.lower(), func.lower(Department.code) == code.lower()),
            Department.id != department.id).first()
        if duplicate:
            flash("A department with that name or code already exists.", "danger")
        else:
            department.name, department.code = name, code
            audit("department_updated", "department", department.id, code)
            db.session.commit()
            flash("Department updated.", "success")
            return redirect(url_for("main.departments"))
    return render_template("department_edit.html", form=form, department=department)


@bp.route("/admin/classes", methods=["GET", "POST"])
@role_required("admin")
def classes():
    form = SubjectForm()
    form.department_id.choices = [(d.id, f"{d.name} ({d.code})") for d in Department.query.order_by(Department.name)]
    form.teacher_id.choices = [(0, "Unassigned")] + [(u.id, u.name) for u in User.query.filter_by(role="teacher", status="active").order_by(User.name)]
    form.student_ids.choices = [(u.id, f"{u.name} · {u.email}") for u in User.query.filter_by(role="student", status="active").order_by(User.name)]
    if form.validate_on_submit():
        code = form.code.data.strip().upper()
        assigned_teacher = db.session.get(User, form.teacher_id.data) if form.teacher_id.data else None
        selected_student_ids = set(form.student_ids.data or [])
        valid_student_ids = valid_batch_student_ids(selected_student_ids, form.department_id.data, form.study_year.data)
        if Subject.query.filter(func.lower(Subject.code) == code.lower()).first():
            flash("That class code already exists.", "danger")
        elif assigned_teacher and assigned_teacher.department_id != form.department_id.data:
            flash("Assign the HOD who belongs to this class's department.", "danger")
        elif valid_student_ids != selected_student_ids:
            flash("Enroll only active students from the selected department and year.", "danger")
        elif form.start_time.data >= form.end_time.data:
            flash("Class end time must be after its start time.", "danger")
        else:
            subject = Subject(name=form.name.data.strip(), code=code,
                              department_id=form.department_id.data,
                              teacher_id=form.teacher_id.data or None,
                              study_year=form.study_year.data, weekday=form.weekday.data,
                              start_time=form.start_time.data, end_time=form.end_time.data,
                              room=form.room.data.strip() or None)
            db.session.add(subject)
            db.session.flush()
            for student_id in selected_student_ids:
                db.session.add(Enrollment(subject_id=subject.id, student_id=student_id))
            audit("class_created", "subject", subject.id, code)
            db.session.commit()
            flash("Class created and enrollment saved.", "success")
            return redirect(url_for("main.classes"))
    rows = Subject.query.order_by(Subject.department_id, Subject.name).all()
    return render_template("classes.html", form=form, subjects=rows)


@bp.route("/admin/classes/<int:subject_id>/edit", methods=["GET", "POST"])
@role_required("admin")
def edit_class(subject_id):
    subject = db.get_or_404(Subject, subject_id)
    form = SubjectForm(obj=subject)
    form.department_id.choices = [(d.id, f"{d.name} ({d.code})") for d in Department.query.order_by(Department.name)]
    form.teacher_id.choices = [(0, "Unassigned")] + [(u.id, u.name) for u in User.query.filter_by(role="teacher", status="active").order_by(User.name)]
    students = User.query.filter_by(role="student", status="active").order_by(User.name).all()
    form.student_ids.choices = [(u.id, f"{u.name} · {u.email}") for u in students]
    if request.method == "GET":
        form.teacher_id.data = subject.teacher_id or 0
        form.student_ids.data = [e.student_id for e in subject.enrollments]
    if form.validate_on_submit():
        code = form.code.data.strip().upper()
        duplicate = Subject.query.filter(func.lower(Subject.code) == code.lower(), Subject.id != subject.id).first()
        assigned_teacher = db.session.get(User, form.teacher_id.data) if form.teacher_id.data else None
        selected_student_ids = set(form.student_ids.data or [])
        valid_student_ids = valid_batch_student_ids(selected_student_ids, form.department_id.data, form.study_year.data)
        if duplicate:
            flash("That class code already exists.", "danger")
        elif assigned_teacher and assigned_teacher.department_id != form.department_id.data:
            flash("Assign the HOD who belongs to this class's department.", "danger")
        elif valid_student_ids != selected_student_ids:
            flash("Enroll only active students from the selected department and year.", "danger")
        elif form.start_time.data >= form.end_time.data:
            flash("Class end time must be after its start time.", "danger")
        else:
            subject.name = form.name.data.strip()
            subject.code = code
            subject.department_id = form.department_id.data
            subject.teacher_id = form.teacher_id.data or None
            subject.study_year = form.study_year.data
            subject.weekday = form.weekday.data
            subject.start_time = form.start_time.data
            subject.end_time = form.end_time.data
            subject.room = form.room.data.strip() or None
            subject.enrollments.clear()
            for student_id in selected_student_ids:
                subject.enrollments.append(Enrollment(student_id=student_id))
            audit("class_updated", "subject", subject.id, code)
            db.session.commit()
            flash("Class details and enrollment updated.", "success")
            return redirect(url_for("main.classes"))
    return render_template("class_edit.html", form=form, subject=subject)


@bp.route("/admin/users")
@role_required("admin")
def users():
    status_filter = request.args.get("status", "")
    query = User.query
    if status_filter in {"active", "pending", "inactive"}:
        query = query.filter_by(status=status_filter)
    rows = query.order_by(User.created_at.desc()).all()
    return render_template("users.html", users=rows, status_filter=status_filter)


@bp.route("/admin/users/new", methods=["GET", "POST"])
@role_required("admin")
def create_user():
    form = AdminCreateUserForm()
    form.department_id.choices = [(0, "No department")] + [(d.id, f"{d.name} ({d.code})") for d in Department.query.order_by(Department.name)]
    if form.validate_on_submit():
        email = form.email.data.strip().lower()
        if User.query.filter(func.lower(User.email) == email).first():
            flash("An account with that email already exists.", "danger")
        elif form.role.data == "admin" and User.query.filter_by(role="admin").first():
            flash("Only one administrator account is allowed. Use the existing admin login.", "danger")
        elif form.role.data in {"teacher", "student"} and not form.department_id.data:
            flash("Select a department for each HOD or student record.", "danger")
        elif form.role.data == "student" and form.study_year.data not in {1, 2, 3}:
            flash("Select the student's year batch.", "danger")
        elif (form.role.data == "student" and
              form.status.data in {"active", "pending"} and
              User.query.filter_by(role="student", department_id=form.department_id.data)
              .filter(User.status.in_(["active", "pending"])).count() >= MAX_STUDENTS_PER_DEPARTMENT):
            flash(f"A department can have at most {MAX_STUDENTS_PER_DEPARTMENT} student records.", "danger")
        else:
            user = User(name=form.name.data.strip(), email=email, role=form.role.data,
                        status=form.status.data, department_id=form.department_id.data or None,
                        study_year=form.study_year.data if form.role.data == "student" else None)
            user.set_password(form.password.data)
            db.session.add(user)
            db.session.flush()
            audit("user_created", "user", user.id, f"role={user.role}; status={user.status}")
            db.session.commit()
            flash("Account created.", "success")
            return redirect(url_for("main.users"))
    return render_template("user_create.html", form=form)


@bp.route("/admin/users/<int:user_id>/edit", methods=["GET", "POST"])
@role_required("admin")
def edit_user(user_id):
    user = db.get_or_404(User, user_id)
    form = UserAdminForm(obj=user)
    form.department_id.choices = [(0, "No department")] + [(d.id, f"{d.name} ({d.code})") for d in Department.query.order_by(Department.name)]
    if request.method == "GET":
        form.department_id.data = user.department_id or 0
        form.study_year.data = user.study_year or 0
    if form.validate_on_submit():
        if form.role.data == "admin" and user.role != "admin" and User.query.filter_by(role="admin").first():
            flash("Only one administrator account is allowed.", "danger")
        elif form.role.data in {"teacher", "student"} and not form.department_id.data:
            flash("Select a department for each HOD or student record.", "danger")
        elif form.role.data == "student" and form.study_year.data not in {1, 2, 3}:
            flash("Select the student's year batch.", "danger")
        elif (form.role.data == "student" and form.department_id.data and
              form.status.data in {"active", "pending"} and
              User.query.filter_by(role="student", department_id=form.department_id.data)
              .filter(User.status.in_(["active", "pending"]), User.id != user.id).count() >= MAX_STUDENTS_PER_DEPARTMENT):
            flash(f"A department can have at most {MAX_STUDENTS_PER_DEPARTMENT} student records.", "danger")
        elif user.id == current_user.id and (form.status.data != "active" or form.role.data != "admin"):
            flash("You cannot deactivate or remove administrator access from your own account.", "danger")
        else:
            user.role = form.role.data
            user.status = form.status.data
            user.department_id = form.department_id.data or None
            user.study_year = form.study_year.data if form.role.data == "student" else None
            details = f"role={user.role}; status={user.status}"
            if form.new_password.data:
                user.set_password(form.new_password.data)
                details += "; administrator set a new password"
            audit("user_updated", "user", user.id, details)
            db.session.commit()
            flash("Account updated.", "success")
            return redirect(url_for("main.users"))
    return render_template("user_edit.html", form=form, user=user)


@bp.route("/attendance/mark", methods=["GET", "POST"])
@role_required("admin", "teacher")
def mark_attendance():
    form = AttendanceSelectorForm()
    subjects = user_subjects(current_user)
    form.subject_id.choices = [(subject.id, f"{subject.code} · {subject.name}") for subject in subjects]
    if not subjects:
        flash("No classes are assigned to your account.", "info")
        return render_template("attendance.html", form=form, subjects=[], roster=[], record=None, selected_subject=None, selected_date=date.today(), selected_session="Regular", item_map={})
    query_subject_id = parse_int(request.values.get("subject_id")) or subjects[0].id
    selected_subject = next((subject for subject in subjects if subject.id == query_subject_id), subjects[0])
    raw_date = request.values.get("date")
    selected_date = parse_date(raw_date) or date.today()
    selected_session = (request.values.get("session") or "Regular").strip()[:80] or "Regular"
    form.subject_id.data = selected_subject.id
    form.date.data = selected_date
    form.session.data = selected_session
    if request.method == "POST" and form.validate_on_submit():
        selected_subject = next((subject for subject in subjects if subject.id == form.subject_id.data), None)
        if not selected_subject:
            abort(403)
        selected_date, selected_session = form.date.data, form.session.data.strip()
        if request.form.get("intent") == "save":
            loaded_selection = [request.form.get("loaded_subject", ""), request.form.get("loaded_date", ""), request.form.get("loaded_session", "")]
            posted_selection = [str(selected_subject.id), selected_date.isoformat(), selected_session]
            if loaded_selection != posted_selection:
                flash("Your class, date, or session changed. Load the roster again before saving.", "warning")
                return redirect(url_for("main.mark_attendance", subject_id=selected_subject.id,
                                        date=selected_date.isoformat(), session=selected_session))
        record = AttendanceRecord.query.filter_by(subject_id=selected_subject.id, date=selected_date, session=selected_session).first()
        enrollments = Enrollment.query.join(User, Enrollment.student_id == User.id).filter(
            Enrollment.subject_id == selected_subject.id, User.status == "active").all()
        existing_items = {item.student_id: item for item in record.items} if record else {}
        roster_ids = {e.student_id for e in enrollments} | set(existing_items)
        if request.form.get("intent") == "save":
            if not record:
                record = AttendanceRecord(subject_id=selected_subject.id, date=selected_date,
                                          session=selected_session, marked_by_id=current_user.id)
                db.session.add(record)
                db.session.flush()
            else:
                record.marked_by_id = current_user.id
            for student_id in roster_ids:
                status = request.form.get(f"status_{student_id}", "present").lower()
                if status not in ATTENDANCE_STATUSES:
                    status = "present"
                remarks = request.form.get(f"remarks_{student_id}", "").strip()[:500]
                item = existing_items.get(student_id)
                if item:
                    item.status, item.remarks = status, remarks
                else:
                    db.session.add(AttendanceItem(record_id=record.id, student_id=student_id, status=status, remarks=remarks))
            audit("attendance_saved", "attendance_record", record.id,
                  f"{selected_subject.code} · {selected_date} · {selected_session} · {len(roster_ids)} students")
            db.session.commit()
            flash("Attendance saved.", "success")
            return redirect(url_for("main.mark_attendance", subject_id=selected_subject.id,
                                    date=selected_date.isoformat(), session=selected_session))
        return redirect(url_for("main.mark_attendance", subject_id=selected_subject.id,
                                date=selected_date.isoformat(), session=selected_session))
    record = AttendanceRecord.query.filter_by(subject_id=selected_subject.id, date=selected_date, session=selected_session).first()
    enrollments = Enrollment.query.join(User, Enrollment.student_id == User.id).filter(
        Enrollment.subject_id == selected_subject.id, User.status == "active").all()
    students = {e.student_id: e.student for e in enrollments}
    item_map = {item.student_id: item for item in record.items} if record else {}
    for item in (record.items if record else []):
        students.setdefault(item.student_id, item.student)
    roster = sorted(students.values(), key=lambda u: u.name.casefold())
    return render_template("attendance.html", form=form, subjects=subjects, roster=roster,
                           record=record, selected_subject=selected_subject, selected_date=selected_date,
                           selected_session=selected_session, item_map=item_map)


@bp.route("/history")
@role_required("student")
def history():
    items = (AttendanceItem.query.join(AttendanceRecord).join(Subject)
             .filter(AttendanceItem.student_id == current_user.id)
             .order_by(AttendanceRecord.date.desc(), Subject.name).all())
    return render_template("history.html", items=items, threshold=get_threshold())


@bp.route("/schedule")
@role_required("student")
def student_schedule():
    subjects = user_subjects(current_user)
    weekdays = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday")
    subjects.sort(key=lambda subject: (
        weekdays.index(subject.weekday) if subject.weekday in weekdays else len(weekdays),
        subject.start_time or "", subject.name.casefold()))
    return render_template("schedule.html", subjects=subjects)


@bp.route("/leave", methods=["GET", "POST"])
@role_required("student")
def leave_requests():
    form = LeaveRequestForm()
    if form.validate_on_submit():
        uploaded = form.attachment.data
        filename = original_name = None
        if uploaded and uploaded.filename:
            original_name = secure_filename(uploaded.filename)[:255]
            ext = original_name.rsplit(".", 1)[-1].lower() if "." in original_name else ""
            if ext not in ALLOWED_UPLOAD_EXTENSIONS:
                flash("Attach a PDF, image, or Word document.", "danger")
                return render_template("leave.html", form=form, requests=LeaveRequest.query.filter_by(student_id=current_user.id).order_by(LeaveRequest.created_at.desc()).all())
            filename = f"{uuid.uuid4().hex}.{ext}"
            uploaded.save(os.path.join(current_app.config["UPLOAD_FOLDER"], filename))
        leave = LeaveRequest(student_id=current_user.id, start_date=form.start_date.data,
                             end_date=form.end_date.data, reason=form.reason.data.strip(),
                             attachment_filename=filename, attachment_original_name=original_name)
        db.session.add(leave)
        db.session.flush()
        audit("leave_submitted", "leave_request", leave.id)
        db.session.commit()
        flash("Leave request submitted.", "success")
        return redirect(url_for("main.leave_requests"))
    requests = LeaveRequest.query.filter_by(student_id=current_user.id).order_by(LeaveRequest.created_at.desc()).all()
    return render_template("leave.html", form=form, requests=requests)


@bp.route("/leave/review")
@role_required("admin", "teacher")
def review_leaves():
    query = LeaveRequest.query.join(User, LeaveRequest.student_id == User.id)
    if current_user.role == "teacher":
        if current_user.department_id:
            query = query.filter(User.department_id == current_user.department_id, User.role == "student")
        else:
            subject_ids = [subject.id for subject in user_subjects(current_user)]
            student_ids = [row[0] for row in db.session.query(Enrollment.student_id).filter(Enrollment.subject_id.in_(subject_ids)).distinct().all()] if subject_ids else []
            query = query.filter(LeaveRequest.student_id.in_(student_ids))
    status_filter = request.args.get("status", "pending")
    if status_filter in {"pending", "approved", "rejected"}:
        query = query.filter(LeaveRequest.status == status_filter)
    requests = query.order_by(LeaveRequest.created_at.desc()).all()
    return render_template("review_leaves.html", requests=requests, status_filter=status_filter)


@bp.route("/leave/<int:leave_id>/review", methods=["POST"])
@role_required("admin", "teacher")
def decide_leave(leave_id):
    leave = db.get_or_404(LeaveRequest, leave_id)
    if current_user.role == "teacher" and not teacher_can_review(current_user, leave.student_id):
        abort(403)
    form = ReviewLeaveForm()
    if form.validate_on_submit():
        leave.status = form.status.data
        leave.reviewed_by_id = current_user.id
        leave.reviewed_at = utcnow()
        audit(f"leave_{leave.status}", "leave_request", leave.id)
        db.session.commit()
        flash(f"Leave request {leave.status}.", "success")
    else:
        flash("Could not update that request.", "danger")
    return redirect(url_for("main.review_leaves"))


@bp.route("/leave/<int:leave_id>/attachment")
@login_required
def leave_attachment(leave_id):
    leave = db.get_or_404(LeaveRequest, leave_id)
    allowed = current_user.role == "admin" or leave.student_id == current_user.id
    if current_user.role == "teacher":
        allowed = teacher_can_review(current_user, leave.student_id)
    if not allowed:
        abort(403)
    if not leave.attachment_filename:
        abort(404)
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], leave.attachment_filename,
                               as_attachment=True, download_name=leave.attachment_original_name or "attachment")


@bp.route("/reports")
@role_required("admin", "teacher")
def reports():
    query, filters = report_query()
    rows = query.limit(3000).all()
    subjects = user_subjects(current_user)
    if current_user.role == "admin":
        departments = Department.query.order_by(Department.name).all()
        students = User.query.filter_by(role="student", status="active").order_by(User.name).all()
    else:
        visible_department_ids = {subject.department_id for subject in subjects}
        departments = (Department.query.filter(Department.id.in_(visible_department_ids)).order_by(Department.name).all()
                       if visible_department_ids else [])
        subject_ids = [subject.id for subject in subjects]
        student_ids = db.session.query(Enrollment.student_id).filter(Enrollment.subject_id.in_(subject_ids)).distinct().subquery() if subject_ids else None
        students = (User.query.filter(User.id.in_(student_ids)).order_by(User.name).all()
                    if student_ids is not None else [])
    summary = {status: sum(item.status == status for item in rows) for status in ATTENDANCE_STATUSES}
    return render_template("reports.html", items=rows, filters=filters, departments=departments,
                           subjects=subjects, students=students, summary=summary)


@bp.route("/reports/export.csv")
@role_required("admin", "teacher")
def export_csv():
    query, _filters = report_query()
    rows = query.limit(100000).all()
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["Date", "Session", "Department", "Class code", "Class", "Student", "Email", "Status", "Remarks", "Marked by"])
    for item in rows:
        record, subject, student = item.record, item.record.subject, item.student
        writer.writerow([safe_spreadsheet_cell(value) for value in [record.date.isoformat(), record.session, subject.department.name, subject.code,
                         subject.name, student.name, student.email, item.status, item.remarks,
                         record.marked_by.name if record.marked_by else ""]])
    response = make_response("\ufeff" + output.getvalue())
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = "attachment; filename=attendance-report.csv"
    return response


@bp.route("/reports/export.xlsx")
@role_required("admin", "teacher")
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    query, _filters = report_query()
    rows = query.limit(100000).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    headers = ["Date", "Session", "Department", "Class code", "Class", "Student", "Email", "Status", "Remarks", "Marked by"]
    sheet.append(headers)
    for item in rows:
        record, subject, student = item.record, item.record.subject, item.student
        values = [record.date.isoformat(), record.session, subject.department.name, subject.code,
                  subject.name, student.name, student.email, item.status, item.remarks,
                  record.marked_by.name if record.marked_by else ""]
        sheet.append([safe_spreadsheet_cell(value) for value in values])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="142D4E")
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 14, "B": 16, "C": 22, "D": 14, "E": 26, "F": 24, "G": 30, "H": 14, "I": 36, "J": 24}.items():
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    workbook.save(output)
    response = make_response(output.getvalue())
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response.headers["Content-Disposition"] = "attachment; filename=attendance-report.xlsx"
    return response


@bp.route("/reports/export.pdf")
@role_required("admin", "teacher")
def export_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    query, _filters = report_query()
    rows = query.limit(5000).all()
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter), title="Attendance report")
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("ReportCell", parent=styles["BodyText"], fontName="Helvetica", fontSize=7, leading=9, spaceAfter=0)
    content = [Paragraph("Attendance report", styles["Title"]),
               Paragraph(f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]), Spacer(1, 12)]
    data = [["Date", "Session", "Class", "Student", "Status", "Remarks"]]
    def pdf_cell(value):
        text = str(value or "").encode("latin-1", errors="replace").decode("latin-1")
        return Paragraph(escape(text), cell_style)
    for item in rows:
        data.append([pdf_cell(item.record.date.isoformat()), pdf_cell(item.record.session),
                     pdf_cell(f"{item.record.subject.code} - {item.record.subject.name}"), pdf_cell(item.student.name),
                     pdf_cell(item.status.title()), pdf_cell(item.remarks)])
    table = Table(data, repeatRows=1, colWidths=[68, 70, 150, 120, 60, 170])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#142d4e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d9e1ea")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fa")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    content.append(table)
    doc.build(content)
    response = make_response(output.getvalue())
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = "attachment; filename=attendance-report.pdf"
    return response


@bp.route("/admin/settings", methods=["GET", "POST"])
@role_required("admin")
def settings():
    form = SettingsForm()
    if request.method == "GET":
        form.attendance_threshold.data = setting_value("attendance_threshold", "75")
    if form.validate_on_submit():
        setting = db.session.get(SystemSetting, "attendance_threshold")
        if setting is None:
            setting = SystemSetting(key="attendance_threshold", value=form.attendance_threshold.data,
                                    updated_by_id=current_user.id)
            db.session.add(setting)
        else:
            setting.value = form.attendance_threshold.data
            setting.updated_by_id = current_user.id
        audit("settings_updated", "system_setting", "attendance_threshold", form.attendance_threshold.data)
        db.session.commit()
        flash("Settings saved.", "success")
        return redirect(url_for("main.settings"))
    return render_template("settings.html", form=form)


@bp.route("/admin/audit")
@role_required("admin")
def audit_logs():
    logs = AuditLog.query.order_by(AuditLog.created_at.desc()).limit(300).all()
    return render_template("audit.html", logs=logs)


@bp.errorhandler(403)
def forbidden(_error):
    return render_template("errors/403.html"), 403
